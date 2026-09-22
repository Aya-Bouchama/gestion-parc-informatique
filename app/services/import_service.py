"""Import / export des données depuis un fichier Excel (.xlsx) ou CSV.

Fichier attendu
---------------
Feuille « Materiels » (ou première feuille / fichier CSV), une ligne par matériel :
    type | numero_serie | modele | emplacement | id_gresa | date_distribution | nombre

Feuille « Pannes » (facultative) — historique des pannes pour le modèle d'IA :
    numero_serie | date_panne | description

Les noms de colonnes sont tolérants : « N° de série », « Numéro de série »,
« Code GRESA », « Quantité »... sont reconnus (majuscules/accents ignorés).

Règles
------
* numéro de série inconnu  -> création du matériel
* numéro de série existant -> mise à jour ; si l'établissement ou la date change,
  une nouvelle distribution est ajoutée à l'historique (rien n'est écrasé)
* une ligne invalide est signalée (numéro de ligne + erreur) sans bloquer les autres
* mode « simulation » : tout est vérifié, rien n'est enregistré
"""
from __future__ import annotations

import io
import unicodedata
from datetime import date, datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from ..extensions import db
from ..models import Etablissement, Materiel, Panne
from .materiel_service import ValidationError, parse_date, validate_payload

EXTENSIONS = {"xlsx", "xlsm", "xls", "csv"}

COLONNES_MATERIEL = ["type", "numero_serie", "modele", "emplacement",
                     "id_gresa", "date_distribution", "nobre_materiel"]
ENTETES_MATERIEL = ["Type", "Numéro de série", "Modèle", "Emplacement",
                    "ID GRESA", "Date de distribution", "Nombre de matériel"]
COLONNES_PANNE = ["numero_serie", "date_panne", "description"]
ENTETES_PANNE = ["Numéro de série", "Date de la panne", "Description"]

SYNONYMES = {
    "type": ["type", "type materiel", "type de materiel", "categorie", "materiel"],
    "numero_serie": ["numero serie", "numero de serie", "n serie", "n de serie", "no serie",
                     "num serie", "serie", "sn", "s n", "serial", "serial number"],
    "modele": ["modele", "marque", "marque modele", "marque et modele", "reference"],
    "emplacement": ["emplacement", "localisation", "lieu", "salle", "service"],
    "id_gresa": ["id gresa", "gresa", "code gresa", "id_gresa", "etablissement", "code etablissement"],
    "date_distribution": ["date distribution", "date de distribution", "date", "date affectation"],
    "nobre_materiel": ["nombre", "nombre materiel", "nombre de materiel", "nobre materiel",
                       "quantite", "qte", "nb"],
    "date_panne": ["date panne", "date de la panne", "date de panne", "date"],
    "description": ["description", "panne", "detail", "commentaire", "observation"],
}


# --------------------------------------------------------------------------- #
# Lecture
# --------------------------------------------------------------------------- #
def _cle(txt) -> str:
    txt = unicodedata.normalize("NFKD", str(txt)).encode("ascii", "ignore").decode().lower()
    for c in "_-°./'()":
        txt = txt.replace(c, " ")
    return " ".join(txt.split())


def _renommer(df: pd.DataFrame, attendues: list[str]) -> pd.DataFrame:
    mapping = {}
    for col in df.columns:
        k = _cle(col)
        for cible in attendues:
            if cible not in mapping.values() and (k == _cle(cible) or k in SYNONYMES.get(cible, [])):
                mapping[col] = cible
                break
    return df.rename(columns=mapping)


def lire_fichier(contenu: bytes, nom: str) -> tuple[pd.DataFrame, pd.DataFrame | None]:
    ext = nom.rsplit(".", 1)[-1].lower() if "." in nom else ""
    if ext not in EXTENSIONS:
        raise ValidationError({"fichier": "Format accepté : Excel (.xlsx) ou CSV"})
    try:
        if ext == "csv":
            texte = contenu.decode("utf-8-sig", errors="replace")
            sep = ";" if texte.split("\n", 1)[0].count(";") >= texte.split("\n", 1)[0].count(",") else ","
            materiels = pd.read_csv(io.StringIO(texte), sep=sep, dtype=str, keep_default_na=False)
            pannes = None
        else:
            feuilles = pd.read_excel(io.BytesIO(contenu), sheet_name=None, dtype=object)
            noms = {_cle(k): k for k in feuilles}
            materiels = feuilles[noms.get("materiels", next(iter(feuilles)))]
            pannes = feuilles[noms["pannes"]] if "pannes" in noms else None
    except Exception as exc:
        raise ValidationError({"fichier": f"Fichier illisible : {exc}"})

    materiels = _renommer(materiels, COLONNES_MATERIEL)
    manquantes = [c for c in ("type", "numero_serie") if c not in materiels.columns]
    if manquantes:
        raise ValidationError({"fichier": "Colonnes introuvables : " + ", ".join(manquantes)
                               + ". Utilisez le modèle à télécharger."})
    if pannes is not None:
        pannes = _renommer(pannes, COLONNES_PANNE)
    return materiels, pannes


def _valeur(v):
    if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT:
        return None
    if isinstance(v, (datetime, pd.Timestamp)):
        return v.date()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))  # numéros de série / GRESA lus comme nombres par Excel
    s = str(v).strip()
    return s or None


def _ligne_vide(row) -> bool:
    return all(_valeur(v) is None for v in row.values)


# --------------------------------------------------------------------------- #
# Import
# --------------------------------------------------------------------------- #
def importer(contenu: bytes, nom: str, simulation: bool = False) -> dict:
    materiels_df, pannes_df = lire_fichier(contenu, nom)
    rapport = {"simulation": simulation, "lignes": 0, "crees": 0, "mis_a_jour": 0,
               "distributions_ajoutees": 0, "pannes_ajoutees": 0, "erreurs": []}
    vus: set[str] = set()

    for i, row in materiels_df.iterrows():
        ligne = i + 2  # ligne 1 = en-têtes (numérotation Excel)
        if _ligne_vide(row):
            continue
        rapport["lignes"] += 1
        data = {c: _valeur(row.get(c)) for c in COLONNES_MATERIEL}
        data["nobre_materiel"] = data["nobre_materiel"] or 1
        try:
            if data["numero_serie"] in vus:
                raise ValidationError({"numero_serie": "En double dans le fichier"})
            existant = Materiel.query.filter_by(numero_serie=data["numero_serie"]).first() \
                if data["numero_serie"] else None
            clean = validate_payload(data, partial=existant is not None)
            vus.add(clean.get("numero_serie", data["numero_serie"]))

            if existant is None:
                m = Materiel(**{k: clean[k] for k in ("type", "numero_serie", "modele", "emplacement")})
                m.distributions.append(Etablissement(
                    id_gresa=clean["id_gresa"], date_distribution=clean["date_distribution"],
                    nobre_materiel=clean["nobre_materiel"]))
                db.session.add(m)
                rapport["crees"] += 1
            else:
                for k in ("type", "modele", "emplacement"):
                    if k in clean:
                        setattr(existant, k, clean[k])
                d = existant.distribution_courante
                nouvelle = ("id_gresa" in clean and "date_distribution" in clean and
                            (d is None or (d.id_gresa, d.date_distribution) !=
                             (clean["id_gresa"], clean["date_distribution"])))
                if nouvelle:
                    existant.distributions.append(Etablissement(
                        id_gresa=clean["id_gresa"], date_distribution=clean["date_distribution"],
                        nobre_materiel=clean.get("nobre_materiel", 1)))
                    rapport["distributions_ajoutees"] += 1
                elif d is not None and "nobre_materiel" in clean:
                    d.nobre_materiel = clean["nobre_materiel"]
                rapport["mis_a_jour"] += 1
            db.session.flush()
        except ValidationError as exc:
            rapport["erreurs"].append({"feuille": "Materiels", "ligne": ligne, "erreurs": exc.errors})

    if pannes_df is not None:
        for i, row in pannes_df.iterrows():
            ligne = i + 2
            if _ligne_vide(row):
                continue
            serie = _valeur(row.get("numero_serie"))
            try:
                m = Materiel.query.filter_by(numero_serie=serie).first() if serie else None
                if m is None:
                    raise ValidationError({"numero_serie": f"Matériel « {serie} » introuvable"})
                when = parse_date(_valeur(row.get("date_panne")) or "")
                if when > date.today():
                    raise ValidationError({"date_panne": "Date dans le futur"})
                if any(p.date_panne == when for p in m.pannes):
                    continue  # déjà importée : on n'ajoute pas de doublon
                m.pannes.append(Panne(date_panne=when,
                                      description=(_valeur(row.get("description")) or "")[:255]))
                db.session.flush()
                rapport["pannes_ajoutees"] += 1
            except ValueError as exc:
                rapport["erreurs"].append({"feuille": "Pannes", "ligne": ligne,
                                           "erreurs": {"date_panne": str(exc)}})
            except ValidationError as exc:
                rapport["erreurs"].append({"feuille": "Pannes", "ligne": ligne, "erreurs": exc.errors})

    if simulation:
        db.session.rollback()
    else:
        db.session.commit()
    return rapport


# --------------------------------------------------------------------------- #
# Modèle vierge et export
# --------------------------------------------------------------------------- #
_ENTETE_FILL = PatternFill("solid", fgColor="0D6EFD")
_ENTETE_FONT = Font(bold=True, color="FFFFFF")


def _feuille(ws, entetes, lignes, largeurs):
    ws.append(entetes)
    for c in ws[1]:
        c.fill, c.font = _ENTETE_FILL, _ENTETE_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for ligne in lignes:
        ws.append(ligne)
    for i, w in enumerate(largeurs):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"


def _format_dates(ws, colonne: str):
    for cell in ws[colonne][1:]:
        cell.number_format = "DD/MM/YYYY"


def generer_modele() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Materiels"
    _feuille(ws, ENTETES_MATERIEL, [
        ["PC", "5CD1234XYZ", "HP ProDesk 400 G7", "Salle informatique 1", "01234A", date(2022, 3, 15), 12],
        ["Imprimante", "VNB3K12345", "HP LaserJet M404", "Secrétariat", "01234A", date(2022, 3, 15), 1],
    ], [16, 22, 24, 24, 14, 20, 18])
    _format_dates(ws, "F")
    dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1",
                        error="Entier positif attendu", showErrorMessage=True)
    ws.add_data_validation(dv)
    dv.add("G2:G5000")
    dvd = DataValidation(type="date", operator="greaterThan", formula1="DATE(1990,1,1)",
                         error="Date attendue (jj/mm/aaaa)", showErrorMessage=True)
    ws.add_data_validation(dvd)
    dvd.add("F2:F5000")

    wp = wb.create_sheet("Pannes")
    _feuille(wp, ENTETES_PANNE, [
        ["VNB3K12345", date(2024, 1, 10), "Bourrage papier récurrent"],
    ], [22, 18, 45])
    _format_dates(wp, "B")

    wi = wb.create_sheet("Instructions")
    for t in [
        "MODE D'EMPLOI",
        "",
        "1. Feuille « Materiels » : une ligne par matériel. Supprimez les 2 lignes d'exemple.",
        "   Obligatoire : Type, Numéro de série, Modèle, Emplacement, ID GRESA, Date de distribution.",
        "   Nombre de matériel : 1 par défaut si vide.",
        "2. Feuille « Pannes » (facultative) : historique des pannes connues.",
        "   Elle permet d'entraîner le modèle de maintenance prédictive sur VOS données.",
        "3. Dates au format jj/mm/aaaa.",
        "4. Un numéro de série déjà présent en base met à jour le matériel ;",
        "   si l'établissement ou la date change, une nouvelle distribution est ajoutée à l'historique.",
        "5. Dans l'application : Importer > Vérifier (aucune écriture) > Importer.",
        "",
        "Astuce : formatez la colonne « Numéro de série » en Texte pour conserver les zéros initiaux.",
    ]:
        wi.append([t])
    wi["A1"].font = Font(bold=True, size=14)
    wi.column_dimensions["A"].width = 100
    for w in (ws, wp):  # colonne numéro de série en texte
        col = "B" if w is ws else "A"
        for cell in w[col]:
            cell.number_format = "@"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def exporter() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Materiels"
    lignes = []
    for m in Materiel.query.order_by(Materiel.materiel_id).all():
        d = m.distribution_courante
        lignes.append([m.type, m.numero_serie, m.modele, m.emplacement,
                       d.id_gresa if d else None, d.date_distribution if d else None,
                       d.nobre_materiel if d else None])
    _feuille(ws, ENTETES_MATERIEL, lignes, [16, 22, 24, 24, 14, 20, 18])
    _format_dates(ws, "F")

    wp = wb.create_sheet("Pannes")
    pannes = [[p.materiel.numero_serie, p.date_panne, p.description]
              for p in Panne.query.order_by(Panne.date_panne).all()]
    _feuille(wp, ENTETES_PANNE, pannes, [22, 18, 45])
    _format_dates(wp, "B")

    wh = wb.create_sheet("Historique distributions")
    hist = [[e.materiel.numero_serie, e.id_gresa, e.date_distribution, e.nobre_materiel]
            for e in Etablissement.query.order_by(Etablissement.materiel_id, Etablissement.date_distribution)]
    _feuille(wh, ["Numéro de série", "ID GRESA", "Date de distribution", "Nombre"], hist, [22, 14, 20, 12])
    _format_dates(wh, "C")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
